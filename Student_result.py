import openpyxl
import os

FILE_NAME = "student_results.xlsx"


def create_file():
    if not os.path.exists(FILE_NAME):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Results"

        ws.append([
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ])

        wb.save(FILE_NAME)


def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    if all(mark >= 35 for mark in marks):
        status = "PASS"

        if percentage >= 90:
            grade = "A+"
        elif percentage >= 80:
            grade = "A"
        elif percentage >= 70:
            grade = "B"
        elif percentage >= 60:
            grade = "C"
        elif percentage >= 50:
            grade = "D"
        else:
            grade = "E"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


def add_student():
    print("\n===== ADD STUDENT RESULT =====")

    roll = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    student_class = input("Enter Class: ")

    marks = []

    for i in range(1, 6):
        mark = float(input("Enter Subject " + str(i) + " Marks: "))
        marks.append(mark)

    total, percentage, grade, status = calculate_result(marks)

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    ws.append([
        roll, name, student_class,
        marks[0], marks[1], marks[2],
        marks[3], marks[4],
        total, percentage, grade, status
    ])

    wb.save(FILE_NAME)

    print("\nResult Saved Successfully!")
    print("Total      :", total)
    print("Percentage :", format(percentage, ".2f") + "%")
    print("Grade      :", grade)
    print("Status     :", status)


def get_result():
    print("\n===== GET STUDENT RESULT =====")

    roll = input("Enter Roll No: ")

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    found = False

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == roll:

            print("\n===== STUDENT RESULT =====")
            print("Roll No    :", row[0])
            print("Name       :", row[1])
            print("Class      :", row[2])
            print("Total      :", row[8])
            print("Percentage :", format(row[9], ".2f") + "%")
            print("Grade      :", row[10])
            print("Status     :", row[11])

            found = True
            break

    if not found:
        print("Student not found!")


def show_all_data():
    print("\n========== ALL STUDENTS ==========")

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        print(
            row[0], "|",
            row[1], "|",
            row[2], "|",
            row[8], "|",
            format(row[9], ".2f"), "|",
            row[10], "|",
            row[11]
        )


def menu():
    create_file()

    while True:
        print("\n==============================")
        print("  STUDENT RESULT MANAGEMENT")
        print("==============================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("==============================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("Program Closed.")
            break

        else:
            print("Invalid Choice!")

menu()